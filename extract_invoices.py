"""Счета из PDF, сканов и фото -> таблица Excel с проверкой сумм.

Запуск:
    python extract_invoices.py <папка_со_счетами> <итог.xlsx>

PDF с текстовым слоем читаются напрямую, сканы и фото распознаются
Tesseract (нужен русский языковой пакет). Все документы обрабатываются
локально, в интернет ничего не отправляется.

Путь к Tesseract на Windows можно задать переменной окружения TESSERACT_CMD,
например: set TESSERACT_CMD=C:\\Program Files\\Tesseract-OCR\\tesseract.exe
"""
import os
import re
import sys
from datetime import date
from pathlib import Path

import cv2
import numpy as np
import pdfplumber
import pypdfium2 as pdfium
import pytesseract
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

if os.environ.get("TESSERACT_CMD"):
    pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}
UNITS = r"шт|пач|уп|усл|пар|бут|кг|л|м|компл|рул"
MONEY = r"\d{1,3}(?:[ \u00a0]\d{3})*[.,]\d{2}"

RE_HEAD = re.compile(
    r"Сч[её]т\s+на\s+оплату\s*(?:№|N[oо°e]?\.?)\s*([A-ZА-ЯЁ0-9][A-ZА-ЯЁ0-9\-/]*)\s+от\s+(\d{1,2})[.,](\d{1,2})[.,](\d{4})",
    re.IGNORECASE,
)
RE_SUPPLIER = re.compile(r"Поставщик:?\s*(.+?),\s*ИНН\s*(\d{10,12})")
RE_BUYER = re.compile(r"Покупатель:?\s*(.+?),\s*ИНН\s*(\d{10,12})")
RE_ITEM = re.compile(
    rf"^(\d{{1,3}})\s+(.+?)\s+(\d+(?:[.,]\d+)?)\s+({UNITS})\.?\s+({MONEY})\s+({MONEY})$"
)
RE_ITEM_LOOSE = re.compile(rf"^(\d{{1,3}})\s+(.+?)\s+({MONEY})\s+({MONEY})$")
RE_TOTAL_PAY = re.compile(rf"Всего\s+к\s+оплате:?\s*({MONEY})")
UNIT_SET = set(UNITS.split("|"))
DIGIT_LOOKALIKES = str.maketrans({"З": "3", "з": "3", "О": "0", "о": "0", "O": "0", "o": "0",
                                  "l": "1", "I": "1", "|": "1", "Б": "6", "б": "6", "S": "5"})
LATIN_TO_CYR = str.maketrans({"a": "а", "c": "с", "e": "е", "o": "о", "p": "р", "x": "х", "y": "у",
                              "n": "п", "w": "ш", "t": "т", "m": "т", "k": "к", "b": "в", "u": "и",
                              "r": "г", "6": "б"})
RE_TOTAL = re.compile(rf"Итого:?\s*({MONEY})")


def to_number(s):
    return float(s.replace("\u00a0", "").replace(" ", "").replace(",", "."))


def remove_table_lines(gray):
    """Стирает линии таблицы: с ними Tesseract теряет строки с товарами."""
    arr = np.array(gray)
    binary = cv2.threshold(arr, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    h_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (max(40, arr.shape[1] // 30), 1))
    v_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(40, arr.shape[0] // 60)))
    lines = cv2.morphologyEx(binary, cv2.MORPH_OPEN, h_kernel) | cv2.morphologyEx(binary, cv2.MORPH_OPEN, v_kernel)
    lines = cv2.dilate(lines, np.ones((3, 3), np.uint8))
    arr[lines > 0] = 255
    return Image.fromarray(arr)


def ocr_image(img):
    img = ImageOps.grayscale(img)
    if img.width < 2200:
        k = 2200 / img.width
        img = img.resize((int(img.width * k), int(img.height * k)), Image.LANCZOS)
    img = remove_table_lines(ImageOps.autocontrast(img))
    return pytesseract.image_to_string(img, lang="rus+eng", config="--psm 6")


def read_text(path):
    """Возвращает (текст, источник)."""
    if path.suffix.lower() == ".pdf":
        with pdfplumber.open(path) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        if len(text.strip()) > 40:
            return text, "PDF с текстом"
        doc = pdfium.PdfDocument(str(path))
        text = "\n".join(ocr_image(page.render(scale=300 / 72).to_pil()) for page in doc)
        return text, "скан в PDF, распознан"
    return ocr_image(Image.open(path)), "фото/скан, распознан"


def norm_qty(token):
    token = token.translate(DIGIT_LOOKALIKES)
    return to_number(token) if re.fullmatch(r"\d+(?:[.,]\d+)?", token) else None


def edit_distance(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def norm_unit(token):
    token = token.rstrip(".").lower().translate(LATIN_TO_CYR)
    if token in UNIT_SET:
        return token
    close = [u for u in UNIT_SET if len(u) >= 2 and edit_distance(token, u) <= 1]
    return close[0] if len(close) == 1 else None


def fix_name(name):
    """Типичные ошибки распознавания в названиях: латиница вместо кириллицы."""
    name = re.sub(r"(\d)\s+n(?=[,\s]|$)", r"\1 л", name)  # «5 n» -> «5 л»
    words = []
    for w in name.split():
        if re.search(r"[А-Яа-яЁё]", w) and re.search(r"[A-Za-z]", w):
            w = w.translate(LATIN_TO_CYR)
        words.append(w)
    return " ".join(words)


def parse_item(line):
    """Строка товара. Сначала строгий разбор, потом терпимый к ошибкам распознавания."""
    m = RE_ITEM.match(line)
    if m:
        return {"n": int(m.group(1)), "name": fix_name(m.group(2).strip()), "qty": to_number(m.group(3)),
                "unit": m.group(4), "price": to_number(m.group(5)), "sum": to_number(m.group(6)),
                "note": ""}
    m = RE_ITEM_LOOSE.match(line)
    if not m:
        return None
    n, middle, price, total = int(m.group(1)), m.group(2).split(), to_number(m.group(3)), to_number(m.group(4))
    if len(middle) >= 3:
        qty, unit = norm_qty(middle[-2]), norm_unit(middle[-1])
        if qty is not None and unit:
            return {"n": n, "name": fix_name(" ".join(middle[:-2])), "qty": qty, "unit": unit,
                    "price": price, "sum": total, "note": ""}
    if price > 0 and abs(total / price - round(total / price)) < 1e-6 and len(middle) >= 2 and len(middle[-1]) <= 4:
        return {"n": n, "name": fix_name(" ".join(middle[:-1])), "qty": float(round(total / price)), "unit": "",
                "price": price, "sum": total, "note": "кол-во по сумме, ед. не распознана"}
    return None


def clean_line(line):
    line = line.replace("\u00a0", " ")
    line = re.sub(r"[|\[\]{}]", " ", line)
    return re.sub(r"\s+", " ", line).strip()


def parse(text):
    lines = [clean_line(l) for l in text.splitlines()]
    joined = "\n".join(lines)
    res = {"number": None, "date": None, "supplier": None, "supplier_inn": None,
           "buyer": None, "total": None, "items": [], "missing": []}

    m = RE_HEAD.search(joined)
    if m:
        res["number"] = m.group(1)
        try:
            res["date"] = date(int(m.group(4)), int(m.group(3)), int(m.group(2)))
        except ValueError:
            pass
    m = RE_SUPPLIER.search(joined)
    if m:
        res["supplier"], res["supplier_inn"] = m.group(1).strip(), m.group(2)
    m = RE_BUYER.search(joined)
    if m:
        res["buyer"] = m.group(1).strip()
    m = RE_TOTAL_PAY.search(joined) or RE_TOTAL.search(joined)
    if m:
        res["total"] = to_number(m.group(1))

    for line in lines:
        item = parse_item(line)
        if item:
            res["items"].append(item)

    labels = {"number": "номер", "date": "дата", "supplier": "поставщик",
              "supplier_inn": "ИНН", "buyer": "покупатель", "total": "итог"}
    res["missing"] = [v for k, v in labels.items() if res[k] in (None, "")]
    if not res["items"]:
        res["missing"].append("позиции")
    unsure = sum(1 for it in res["items"] if it["note"])
    if unsure:
        res["missing"].append(f"единицы измерения (строк: {unsure})")
    return res


HEAD_FILL = PatternFill("solid", fgColor="1C7A4E")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
THIN = Side(style="thin", color="C9D4CD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CHECK_FILL = PatternFill("solid", fgColor="FFF1B8")
MONEY_FMT = '#,##0.00'


def style_header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def build_workbook(rows, out_path):
    wb = Workbook()
    inv = wb.active
    inv.title = "Счета"
    pos = wb.create_sheet("Позиции")
    about = wb.create_sheet("О файле")

    style_header(inv, ["Файл", "Источник", "№ счёта", "Дата", "Поставщик", "ИНН поставщика",
                       "Покупатель", "Позиций", "Итого по счёту, ₽", "Сумма позиций, ₽",
                       "Расхождение, ₽", "Проверка", "Что не распознано"],
                 [30, 22, 11, 11, 30, 15, 22, 9, 15, 15, 13, 11, 22])
    style_header(pos, ["Файл", "№ счёта", "№", "Наименование", "Кол-во", "Ед.", "Цена, ₽",
                       "Сумма в счёте, ₽", "Кол-во × цена, ₽", "Строка сходится", "Примечание"],
                 [30, 11, 5, 40, 8, 6, 12, 15, 15, 14, 32])

    p = 2
    for r, row in enumerate(rows, start=2):
        f, src, d = row["file"], row["source"], row["data"]
        inv.append([f, src, d["number"], d["date"], d["supplier"], d["supplier_inn"], d["buyer"],
                    len(d["items"]), d["total"],
                    f'=SUMIFS(Позиции!$H:$H,Позиции!$A:$A,$A{r})',
                    f'=IF(I{r}="","",ROUND(I{r}-J{r},2))',
                    f'=IF(OR(M{r}<>"",K{r}=""),"проверить",IF(ABS(K{r})<0.01,"ок","проверить"))',
                    ", ".join(d["missing"])])
        for it in d["items"]:
            pos.append([f, d["number"], it["n"], it["name"], it["qty"], it["unit"], it["price"],
                        it["sum"], f"=ROUND(E{p}*G{p},2)",
                        f'=IF(OR(K{p}<>"",ABS(H{p}-I{p})>=0.01),"проверить","ок")', it["note"]])
            p += 1

    last_inv = len(rows) + 1
    total_row = last_inv + 1
    inv.cell(row=total_row, column=1, value="Итого по всем счетам")
    inv.cell(row=total_row, column=9, value=f"=SUM(I2:I{last_inv})")
    inv.cell(row=total_row, column=10, value=f"=SUM(J2:J{last_inv})")
    inv.cell(row=total_row, column=11, value=f"=SUM(K2:K{last_inv})")

    for ws, last, money_cols, date_cols in ((inv, total_row, (9, 10, 11), (4,)), (pos, p - 1, (7, 8, 9), ())):
        for row in ws.iter_rows(min_row=2, max_row=last):
            for c in row:
                c.font, c.border = BODY_FONT, BORDER
                c.alignment = Alignment(vertical="center")
                if c.column in money_cols:
                    c.number_format = MONEY_FMT
                if c.column in date_cols:
                    c.number_format = "DD.MM.YYYY"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{last if ws is pos else last_inv}"
    for c in inv[total_row]:
        c.font = BOLD_FONT

    inv.conditional_formatting.add(f"A2:M{last_inv}", FormulaRule(formula=[f'$L2="проверить"'], fill=CHECK_FILL))
    pos.conditional_formatting.add(f"A2:K{p - 1}", FormulaRule(formula=[f'$J2="проверить"'], fill=CHECK_FILL))

    notes = [
        ("Что это", "Сводка счетов на оплату, собранная скриптом из PDF, сканов и фото."),
        ("Счета", "Одна строка на документ. Столбцы J–L считаются формулами по листу «Позиции»."),
        ("Позиции", "Все строки товаров и услуг из счетов. Столбцы I–J — формулы проверки."),
        ("Проверка «ок»", "Сумма позиций совпала с итогом в документе, все поля распознаны."),
        ("Проверка «проверить»", "Итог в документе не совпал с суммой позиций или поле не распознано. "
                                 "Такие строки подсвечены жёлтым, их стоит сверить с оригиналом."),
    ]
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 90
    for a, b in notes:
        about.append([a, b])
    for row in about.iter_rows():
        row[0].font = BOLD_FONT
        row[1].font = BODY_FONT
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    for ws in (inv, pos):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)


def main(src, out):
    files = sorted(p for p in Path(src).iterdir()
                   if p.suffix.lower() == ".pdf" or p.suffix.lower() in IMAGE_EXT)
    rows = []
    for f in files:
        text, source = read_text(f)
        data = parse(text)
        rows.append({"file": f.name, "source": source, "data": data})
        flag = "" if not data["missing"] else f"  не распознано: {', '.join(data['missing'])}"
        print(f"{f.name}: {source}, позиций {len(data['items'])}, итог {data['total']}{flag}")
    rows.sort(key=lambda r: (r["data"]["date"] or date.max, r["file"]))
    build_workbook(rows, out)
    print(f"\nГотово: {out}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
